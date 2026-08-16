from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from docx import Document

from cover_letter_manager.pdf_convert import find_libreoffice
from cover_letter_manager.repository import CoverLetterRepository, CreateEntryRequest


@pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice is not installed")
def test_docx_is_preserved_and_converted(tmp_path: Path) -> None:
    source = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("Dear Hiring Manager,")
    document.add_paragraph("I build reliable backend APIs and payment services.")
    document.save(source)

    repository = CoverLetterRepository(tmp_path / "library")
    entry = repository.create_entry(
        CreateEntryRequest(
            company="Finance Co",
            position="Backend Developer",
            application_date=date(2026, 8, 15),
            source_file=source,
        )
    )

    assert entry.original_path.suffix == ".docx"
    assert entry.original_path.exists()
    assert entry.pdf_path.exists()
    assert entry.pdf_path.read_bytes().startswith(b"%PDF")
