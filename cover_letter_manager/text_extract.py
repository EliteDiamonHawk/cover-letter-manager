from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".docx",
    ".odt",
    ".rtf",
    ".html",
    ".htm",
}


class TextExtractionError(RuntimeError):
    pass


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip_depth += 1
        elif tag in {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip_depth:
            self._skip_depth -= 1
        elif tag in {"p", "div", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)

    def text(self) -> str:
        return "".join(self.parts)


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    raise TextExtractionError(f"Could not decode text file: {path}")


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_odt(path: Path) -> str:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import H, P

    document = load(str(path))
    parts: list[str] = []
    for node in document.getElementsByType(P) + document.getElementsByType(H):
        parts.append(teletype.extractText(node))
    return "\n".join(parts)


def _extract_rtf(path: Path) -> str:
    raw = _read_text(path)
    try:
        from striprtf.striprtf import rtf_to_text

        return rtf_to_text(raw)
    except ImportError:
        # A deliberately conservative fallback for environments where striprtf
        # has not yet been installed. It is adequate for simple RTF letters.
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
        text = text.replace("{", " ").replace("}", " ")
        return text


def _extract_html(path: Path) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(_read_text(path))
    return parser.text()


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: Path) -> str:
    """Extract searchable text from a supported file."""

    if not path.exists():
        raise TextExtractionError(f"File does not exist: {path}")

    suffix = path.suffix.casefold()
    try:
        if suffix in {".txt", ".md", ".markdown"}:
            text = _read_text(path)
        elif suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".odt":
            text = _extract_odt(path)
        elif suffix == ".rtf":
            text = _extract_rtf(path)
        elif suffix in {".html", ".htm"}:
            text = _extract_html(path)
        else:
            raise TextExtractionError(f"Text extraction is not supported for {suffix or 'this file type'}.")
    except TextExtractionError:
        raise
    except Exception as exc:
        raise TextExtractionError(f"Could not extract text from {path.name}: {exc}") from exc

    return normalize_text(text)
